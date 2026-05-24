from __future__ import annotations

from collections import Counter
from io import BytesIO
from typing import Any, Iterable

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

from app.services.paper_format_known_profiles import detect_known_profile_key
from app.services.paper_format_normalizer import _build_paragraph_contexts
from app.services.paper_format_requirement_rules import extract_requirements_and_notes_from_text


def extract_template_requirements_from_docx(content: bytes) -> dict[str, Any]:
    document = Document(BytesIO(content))
    contexts = _build_paragraph_contexts(document.paragraphs)
    requirements: dict[str, Any] = {}
    notes: list[str] = []

    if document.sections:
        section = document.sections[0]
        requirements["margin_top_cm"] = round(section.top_margin.cm, 2)
        requirements["margin_bottom_cm"] = round(section.bottom_margin.cm, 2)
        requirements["margin_left_cm"] = round(section.left_margin.cm, 2)
        requirements["margin_right_cm"] = round(section.right_margin.cm, 2)
        if _footer_has_page_number(section):
            requirements["add_page_number"] = True

    normal_style = _find_style(document, ("Normal", "\u6b63\u6587"))
    if normal_style is not None:
        _merge_style_requirements(
            requirements,
            normal_style,
            font_field="body_font",
            latin_field="latin_font",
            size_field="body_size_pt",
            spacing_field="line_spacing",
            indent_field="first_line_indent_cm",
            align_field="body_align",
        )

    _merge_heading_style_requirements(document, requirements, 1, "heading_1")
    _merge_heading_style_requirements(document, requirements, 2, "heading_2")
    _merge_heading_style_requirements(document, requirements, 3, "heading_3")

    _merge_paragraph_fallback(document, contexts, requirements, role="body")
    _merge_paragraph_fallback(document, contexts, requirements, role="heading_1")
    _merge_paragraph_fallback(document, contexts, requirements, role="heading_2")
    _merge_paragraph_fallback(document, contexts, requirements, role="heading_3")
    _merge_paragraph_fallback(document, contexts, requirements, role="reference_item")
    _merge_appendix_fallback(document, contexts, requirements)

    table_size = _extract_table_size_pt(document)
    if table_size is not None:
        requirements["table_size_pt"] = table_size

    plain_text = "\n".join(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
    matched_profile_key = detect_known_profile_key(plain_text)
    text_requirements, text_notes = extract_requirements_and_notes_from_text(plain_text)
    requirements.update({key: value for key, value in text_requirements.items() if value is not None})
    for conservative_key in ("latin_font", "reference_hanging_indent_cm", "reference_line_spacing"):
        if conservative_key in requirements and conservative_key not in text_requirements:
            requirements.pop(conservative_key, None)
    notes.extend(text_notes)

    requirements = {key: value for key, value in requirements.items() if value is not None}
    notes = _dedupe_notes(notes)

    if not requirements:
        notes.append("\u6a21\u677f\u4e2d\u6682\u672a\u63d0\u53d6\u5230\u7a33\u5b9a\u53ef\u7528\u7684\u6837\u5f0f\u5b57\u6bb5\u3002")

    return {
        "requirements": requirements,
        "notes": notes,
        "matched_profile_key": matched_profile_key,
        "summary": {
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
            "detected_fields": len(requirements),
        },
    }


def _dedupe_notes(notes: list[str]) -> list[str]:
    deduped: list[str] = []
    seen: set[str] = set()
    for note in notes:
        text = " ".join(str(note or "").split()).strip()
        if not text or text in seen:
            continue
        deduped.append(text)
        seen.add(text)
    return deduped[:8]


def _merge_heading_style_requirements(document: Document, requirements: dict[str, Any], level: int, prefix: str) -> None:
    style = _find_style(
        document,
        (
            f"Heading {level}",
            f"\u6807\u9898 {level}",
            f"\u6807\u9898{level}",
        ),
    )
    if style is None:
        return
    _merge_style_requirements(
        requirements,
        style,
        font_field=f"{prefix}_font",
        latin_field="latin_font",
        size_field=f"{prefix}_size_pt",
        spacing_field=None,
        indent_field=None,
        align_field=f"{prefix}_align",
    )


def _merge_style_requirements(
    requirements: dict[str, Any],
    style,
    *,
    font_field: str | None,
    latin_field: str | None,
    size_field: str | None,
    spacing_field: str | None,
    indent_field: str | None,
    align_field: str | None,
) -> None:
    font_name = _get_style_east_asian_font(style)
    latin_font = _get_style_latin_font(style)
    size_pt = _get_style_size_pt(style)
    spacing = _get_style_line_spacing(style)
    indent_cm = _get_style_first_line_indent_cm(style)
    align = _get_style_align(style)

    if font_field and font_name and font_field not in requirements:
        requirements[font_field] = font_name
    if latin_field and latin_font and latin_field not in requirements:
        requirements[latin_field] = latin_font
    if size_field and size_pt is not None and size_field not in requirements:
        requirements[size_field] = size_pt
    if spacing_field and spacing is not None and spacing_field not in requirements:
        requirements[spacing_field] = spacing
    if indent_field and indent_cm is not None and indent_field not in requirements:
        requirements[indent_field] = indent_cm
    if align_field and align and align_field not in requirements:
        requirements[align_field] = align


def _merge_paragraph_fallback(document: Document, contexts, requirements: dict[str, Any], *, role: str) -> None:
    paragraphs = [paragraph for paragraph, context in zip(document.paragraphs, contexts) if context.role == role and paragraph.text.strip()]
    if not paragraphs:
        return
    sample = paragraphs[0]
    font_name, latin_font, size_pt = _extract_run_style(sample.runs)
    align = _alignment_to_name(sample.alignment)
    spacing = _line_spacing_to_value(sample.paragraph_format.line_spacing, sample.paragraph_format.line_spacing_rule)
    indent_cm = _length_to_cm(sample.paragraph_format.first_line_indent)

    mapping = {
        "body": ("body_font", "latin_font", "body_size_pt", "body_align", "line_spacing", "first_line_indent_cm"),
        "heading_1": ("heading_1_font", "latin_font", "heading_1_size_pt", "heading_1_align", None, None),
        "heading_2": ("heading_2_font", "latin_font", "heading_2_size_pt", "heading_2_align", None, None),
        "heading_3": ("heading_3_font", "latin_font", "heading_3_size_pt", "heading_3_align", None, None),
        "reference_item": ("body_font", "latin_font", "body_size_pt", None, "reference_line_spacing", None),
    }
    fields = mapping.get(role)
    if not fields:
        return
    font_field, latin_field, size_field, align_field, spacing_field, indent_field = fields
    if font_field and font_name and font_field not in requirements:
        requirements[font_field] = font_name
    if latin_field and latin_font and latin_field not in requirements:
        requirements[latin_field] = latin_font
    if size_field and size_pt is not None and size_field not in requirements:
        requirements[size_field] = size_pt
    if align_field and align and align_field not in requirements:
        requirements[align_field] = align
    if spacing_field and spacing is not None and spacing_field not in requirements:
        requirements[spacing_field] = spacing
    if indent_field and indent_cm is not None and indent_field not in requirements:
        requirements[indent_field] = indent_cm

    if role == "reference_item" and "reference_hanging_indent_cm" not in requirements:
        left_indent = _length_to_cm(sample.paragraph_format.left_indent)
        first_indent = _length_to_cm(sample.paragraph_format.first_line_indent)
        if left_indent is not None and first_indent is not None and first_indent < 0:
            requirements["reference_hanging_indent_cm"] = round(left_indent, 2)


def _merge_appendix_fallback(document: Document, contexts, requirements: dict[str, Any]) -> None:
    buckets: dict[str, list] = {"english": [], "translation": [], "code": []}
    for paragraph, context in zip(document.paragraphs, contexts):
        if context.role == "appendix_body" and context.appendix_kind in buckets and paragraph.text.strip():
            buckets[context.appendix_kind].append(paragraph)

    mapping = {
        "english": ("appendix_english_font", "appendix_english_size_pt", "appendix_english_line_spacing"),
        "translation": ("appendix_translation_font", "appendix_translation_size_pt", "appendix_translation_line_spacing"),
        "code": ("appendix_code_font", "appendix_code_size_pt", "appendix_code_line_spacing"),
    }
    for kind, paragraphs in buckets.items():
        if not paragraphs:
            continue
        sample = paragraphs[0]
        font_name, latin_font, size_pt = _extract_run_style(sample.runs)
        spacing = _line_spacing_to_value(sample.paragraph_format.line_spacing, sample.paragraph_format.line_spacing_rule)
        font_field, size_field, spacing_field = mapping[kind]
        chosen_font = latin_font if kind in {"english", "code"} else font_name
        if chosen_font and font_field not in requirements:
            requirements[font_field] = chosen_font
        if size_pt is not None and size_field not in requirements:
            requirements[size_field] = size_pt
        if spacing is not None and spacing_field not in requirements:
            requirements[spacing_field] = spacing


def _find_style(document: Document, names: Iterable[str]):
    candidates = {name.lower() for name in names}
    for style in document.styles:
        if style.type != WD_STYLE_TYPE.PARAGRAPH:
            continue
        if style.name and style.name.lower() in candidates:
            return style
    return None


def _extract_table_size_pt(document: Document) -> float | None:
    sizes: list[float] = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _, _, size_pt = _extract_run_style(paragraph.runs)
                    if size_pt is not None:
                        sizes.append(size_pt)
    return _pick_common_number(sizes)


def _extract_run_style(runs) -> tuple[str | None, str | None, float | None]:
    east_fonts: list[str] = []
    latin_fonts: list[str] = []
    sizes: list[float] = []
    for run in runs:
        east = _get_run_east_asian_font(run)
        latin = getattr(run.font, "name", None)
        size = getattr(getattr(run.font, "size", None), "pt", None)
        if east:
            east_fonts.append(east)
        if latin:
            latin_fonts.append(latin)
        if size:
            sizes.append(round(float(size), 1))
    return _pick_common_text(east_fonts), _pick_common_text(latin_fonts), _pick_common_number(sizes)


def _get_style_east_asian_font(style) -> str | None:
    try:
        r_pr = style.element.get_or_add_rPr()
        r_fonts = r_pr.rFonts
        if r_fonts is not None:
            value = r_fonts.get(qn("w:eastAsia"))
            if value:
                return value
    except Exception:
        return None
    return None


def _get_style_latin_font(style) -> str | None:
    value = getattr(style.font, "name", None)
    return str(value).strip() if value else None


def _get_style_size_pt(style) -> float | None:
    size = getattr(getattr(style.font, "size", None), "pt", None)
    if size is None:
        return None
    return round(float(size), 1)


def _get_style_line_spacing(style) -> float | None:
    paragraph_format = getattr(style, "paragraph_format", None)
    if paragraph_format is None:
        return None
    return _line_spacing_to_value(paragraph_format.line_spacing, paragraph_format.line_spacing_rule)


def _get_style_first_line_indent_cm(style) -> float | None:
    paragraph_format = getattr(style, "paragraph_format", None)
    if paragraph_format is None:
        return None
    return _length_to_cm(paragraph_format.first_line_indent)


def _get_style_align(style) -> str | None:
    paragraph_format = getattr(style, "paragraph_format", None)
    if paragraph_format is None:
        return None
    return _alignment_to_name(paragraph_format.alignment)


def _get_run_east_asian_font(run) -> str | None:
    try:
        r_pr = run._element.rPr
        if r_pr is None or r_pr.rFonts is None:
            return None
        value = r_pr.rFonts.get(qn("w:eastAsia"))
        return str(value).strip() if value else None
    except Exception:
        return None


def _alignment_to_name(value) -> str | None:
    if value == WD_ALIGN_PARAGRAPH.CENTER:
        return "center"
    if value == WD_ALIGN_PARAGRAPH.RIGHT:
        return "right"
    if value == WD_ALIGN_PARAGRAPH.JUSTIFY:
        return "justify"
    if value == WD_ALIGN_PARAGRAPH.LEFT:
        return "left"
    return None


def _line_spacing_to_value(line_spacing, rule) -> float | None:
    if rule == WD_LINE_SPACING.SINGLE:
        return 1.0
    if rule == WD_LINE_SPACING.ONE_POINT_FIVE:
        return 1.5
    if rule == WD_LINE_SPACING.DOUBLE:
        return 2.0
    if isinstance(line_spacing, (int, float)):
        if 0.8 <= float(line_spacing) <= 3.5:
            return round(float(line_spacing), 2)
    return None


def _length_to_cm(length_value) -> float | None:
    if length_value is None:
        return None
    try:
        return round(float(length_value.cm), 2)
    except Exception:
        return None


def _footer_has_page_number(section) -> bool:
    try:
        for paragraph in section.footer.paragraphs:
            xml = paragraph._p.xml
            if "PAGE" in xml:
                return True
    except Exception:
        return False
    return False


def _pick_common_text(values: list[str]) -> str | None:
    if not values:
        return None
    return Counter(values).most_common(1)[0][0]


def _pick_common_number(values: list[float]) -> float | None:
    if not values:
        return None
    return Counter(values).most_common(1)[0][0]
