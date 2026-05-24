from __future__ import annotations

import json
import re
import fitz
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from time import time_ns
from datetime import datetime, timedelta, timezone
from typing import Annotated
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile, status
from fastapi.responses import PlainTextResponse, Response
from sqlalchemy import delete as sql_delete, select, update as sql_update
from sqlalchemy.orm.attributes import flag_modified
from sqlalchemy.orm import Session

from app.api.deps import get_current_user
from app.core.config import settings
from app.db.session import get_db
from app.models import AiProvider, Annotation, Folder, InkAnnotation, Paper, PaperFullTranslation, PaperLiteratureCache, PaperSummary, User
from app.schemas.paper import (
    FolderCreate,
    FullTranslationResponse,
    FullTranslationStartRequest,
    PAPER_TEXT_LIMITS,
    FolderResponse,
    FolderUpdate,
    PaperMetadata,
    PaperResponse,
    PaperTrashResponse,
    PaperUpdate,
    compact_paper_text,
)
from app.services.crypto import decrypt_api_key
from app.services.ai_provider_manager import resolve_user_provider
from app.services.background_task_runner import spawn_full_translation_worker_process
from app.services.notification import compact_notification_text, create_notification
from app.services.membership import ensure_notes_export_allowed
from app.services.paper_metadata import (
    extract_pdf_metadata,
    is_missing_metadata_value,
    is_weak_metadata_title,
)
from app.services.oss_storage import (
    delete_object as delete_oss_object,
    delete_prefix as delete_oss_prefix,
    download_bytes as download_oss_bytes,
    download_file as download_oss_file,
    object_exists as oss_object_exists,
    page_image_object_key,
    page_image_prefix,
    paper_object_key,
    upload_bytes as upload_oss_bytes,
    upload_file as upload_oss_file,
)
from app.services.translate import translate_title
from app.services.upload_mirror import mirror_upload_file, remove_mirrored_upload

router = APIRouter(prefix="/papers", tags=["papers"])

ALLOWED_PDF_TYPES = {"application/pdf"}
TRASH_RETENTION = timedelta(days=7)
UPLOAD_READ_CHUNK_SIZE = 1024 * 1024
FULL_TRANSLATION_DISABLED_MESSAGE = "全文翻译功能已暂停。"


def ensure_full_translation_enabled() -> None:
    if not settings.full_translation_enabled:
        raise HTTPException(status_code=404, detail=FULL_TRANSLATION_DISABLED_MESSAGE)


def build_folder_response(folder: Folder) -> FolderResponse:
    return FolderResponse(
        id=folder.id,
        name=folder.name,
        created_at=folder.created_at.isoformat() if folder.created_at else None,
    )


def _build_public_file_path(file_path: str) -> str:
    if not file_path:
        return file_path
    if file_path.startswith(("http://", "https://")):
        return file_path
    base_url = settings.upload_public_base_url
    if base_url and file_path.startswith("/"):
        return f"{base_url}{file_path}"
    return file_path


def build_paper_response(paper: Paper) -> PaperResponse:
    return PaperResponse(
        id=paper.id,
        folder_id=paper.folder_id,
        file_name=paper.file_name,
        file_path=_build_public_file_path(paper.file_path),
        file_size=paper.file_size,
        title=paper.title or "",
        translated_title=paper.translated_title,
        author=paper.author,
        subject=paper.subject,
        keywords=paper.keywords,
        creator=paper.creator,
        producer=paper.producer,
        creation_date=paper.creation_date,
        modification_date=paper.modification_date,
        doi=paper.doi,
        arxiv_id=paper.arxiv_id,
        page_count=paper.page_count,
        last_viewed_at=paper.last_viewed_at.isoformat() if paper.last_viewed_at else None,
        created_at=paper.created_at.isoformat() if paper.created_at else None,
    )


def build_trash_response(paper: Paper, folder_name: str = "未分类") -> PaperTrashResponse:
    deleted_at = paper.deleted_at or datetime.now(timezone.utc)
    expires_at = deleted_at + TRASH_RETENTION
    return PaperTrashResponse(
        id=paper.id,
        folder_id=paper.deleted_original_folder_id or paper.folder_id,
        folder_name=folder_name,
        file_name=paper.file_name,
        file_size=paper.file_size,
        title=paper.title or paper.file_name,
        author=paper.author,
        page_count=paper.page_count or 0,
        deleted_at=deleted_at.isoformat(),
        expires_at=expires_at.isoformat(),
    )


def active_paper_query(paper_id: int, user_id: int):
    return select(Paper).where(
        Paper.id == paper_id,
        Paper.user_id == user_id,
        Paper.deleted_at.is_(None),
    )


async def _persist_uploaded_pdf(upload: UploadFile, destination: Path, *, max_size_bytes: int) -> int:
    written = 0
    try:
        with destination.open("wb") as handle:
            while True:
                chunk = await upload.read(UPLOAD_READ_CHUNK_SIZE)
                if not chunk:
                    break
                written += len(chunk)
                if written > max_size_bytes:
                    raise HTTPException(
                        status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                        detail=f"PDF file exceeds the {max_size_bytes // (1024 * 1024)} MB upload limit.",
                    )
                handle.write(chunk)
    except Exception:
        try:
            destination.unlink()
        except OSError:
            pass
        raise
    finally:
        await upload.close()
    return written


def _append_translation_debug_log(message: str) -> None:
    if not settings.translation_debug_log_enabled:
        return
    log_path = Path(settings.uploads_dir).resolve().parent / "translate_debug.log"
    log_path.parent.mkdir(parents=True, exist_ok=True)
    with log_path.open("a", encoding="utf-8") as handle:
        handle.write(message)


def _clear_generated_content(db: Session, paper_id: int, user_id: int) -> None:
    db.execute(sql_delete(PaperFullTranslation).where(PaperFullTranslation.paper_id == paper_id))
    db.execute(
        sql_delete(PaperSummary).where(
            PaperSummary.paper_id == paper_id,
            PaperSummary.user_id == user_id,
        )
    )


def _permanently_delete_paper(db: Session, paper: Paper) -> None:
    actual_file = _resolve_paper_file(paper.file_path)
    if actual_file and actual_file.exists():
        try:
            actual_file.unlink()
            remove_mirrored_upload(f"papers/{actual_file.name}")
        except OSError:
            pass
    delete_oss_object(paper_object_key(paper.file_path))
    delete_oss_prefix(page_image_prefix(paper.file_path))
    db.delete(paper)


def _purge_expired_trashed_papers(db: Session, user_id: int | None = None) -> int:
    cutoff = datetime.now(timezone.utc) - TRASH_RETENTION
    query = select(Paper).where(Paper.deleted_at.is_not(None), Paper.deleted_at < cutoff)
    if user_id is not None:
        query = query.where(Paper.user_id == user_id)
    papers = db.scalars(query).all()
    for paper in papers:
        _permanently_delete_paper(db, paper)
    return len(papers)


def build_full_translation_response(item: PaperFullTranslation | None) -> FullTranslationResponse:
    if not item:
        return FullTranslationResponse()
    parse_summary = dict(getattr(item, "parse_summary", None) or {})
    parse_summary["diagnostics"] = build_full_translation_diagnostics(item.pages_json or [])
    pending_blocks_count = sum(
        1
        for page in (item.pages_json or [])
        for block in (page.get("blocks") or [])
        if block.get("status") == "pending"
    )
    failed_blocks_count = sum(
        1
        for page in (item.pages_json or [])
        for block in (page.get("blocks") or [])
        if block.get("status") == "failed"
    )
    return FullTranslationResponse(
        status=item.status if item.status in {"idle", "running", "completed", "partial_failed", "error", "cancelled"} else "idle",
        source_hash=item.source_hash or "",
        pages=item.pages_json or [],
        completed_units=item.completed_units or 0,
        total_units=item.total_units or 0,
        error_message=item.error_message,
        provider_id=item.provider_id,
        parse_mode=getattr(item, "parse_mode", "auto") or "auto",
        parse_engine=getattr(item, "parse_engine", "local") or "local",
        parse_summary=parse_summary,
        translation_engine=getattr(item, "translation_engine", "ai") or "ai",
        termbase_version=getattr(item, "termbase_version", "") or "",
        failed_blocks_count=failed_blocks_count + pending_blocks_count,
        pending_blocks_count=pending_blocks_count,
    )


COPY_BLOCK_TYPES = {"formula", "image", "table", "page_meta"}
COPY_BLOCK_KINDS = {"footer", "caption"}


def normalize_block_type(block: dict) -> str:
    kind = str(block.get("kind") or "paragraph").strip() or "paragraph"
    block_type = str(block.get("type") or "").strip() or "text"
    text = str(block.get("source_text") or "").strip()
    if kind == "caption":
        block_type = "caption"
    if re.search(r"^(fig\.|figure|table)\s*\d+", text, re.I):
        block_type = "caption"
    if re.search(r"[∑∫√≈≤≥∞α-ωΑ-Ω]|\\[a-zA-Z]+|\$[^$]+\$", text):
        block_type = "formula"
    block["type"] = block_type
    return block_type


def should_copy_block(block: dict) -> bool:
    block_type = normalize_block_type(block)
    kind = str(block.get("kind") or "").strip()
    return bool(
        block.get("skip_translate")
        or block_type in COPY_BLOCK_TYPES
        or kind in COPY_BLOCK_KINDS
        or str(block.get("translate_policy") or "") in {"copy", "skip"}
    )


def normalize_translation_pages(pages: list[dict]) -> list[dict]:
    normalized = deepcopy(pages or [])
    for page in normalized:
        for block in page.get("blocks") or []:
            normalize_block_type(block)
            if should_copy_block(block):
                block["translate_policy"] = "copy"
                block["status"] = "copied"
                block["translated_text"] = block.get("source_text", "")
            else:
                block["translate_policy"] = "translate"
                block["status"] = block.get("status") if block.get("status") in {"translated", "failed"} else "pending"
                block["translated_text"] = block.get("translated_text", "")
    return normalized


def count_translatable_units(pages: list[dict]) -> int:
    return sum(
        1
        for page in pages or []
        for block in page.get("blocks") or []
        if not should_copy_block(block) and str(block.get("source_text") or "").strip()
    )


def is_translation_copy_like_block(block: dict) -> bool:
    block_type = str(block.get("type") or "").strip()
    kind = str(block.get("kind") or "").strip()
    policy = str(block.get("translate_policy") or "").strip()
    status = str(block.get("status") or "").strip()
    return bool(
        block.get("skip_translate")
        or block_type in COPY_BLOCK_TYPES
        or kind in COPY_BLOCK_KINDS
        or policy in {"copy", "skip"}
        or status == "copied"
    )


def build_full_translation_diagnostics(pages: list[dict]) -> dict:
    stats = {
        "page_count": len(pages or []),
        "block_count": 0,
        "translatable_blocks": 0,
        "copied_blocks": 0,
        "translated_blocks": 0,
        "failed_blocks": 0,
        "pending_blocks": 0,
        "source_chars": 0,
        "translated_chars": 0,
        "coverage_percent": 0,
        "failure_percent": 0,
    }
    for page in pages or []:
        for block in page.get("blocks") or []:
            stats["block_count"] += 1
            source_text = str(block.get("source_text") or "").strip()
            translated_text = str(block.get("translated_text") or "").strip()
            stats["source_chars"] += len(source_text)
            if is_translation_copy_like_block(block):
                stats["copied_blocks"] += 1
                continue
            if not source_text:
                continue
            stats["translatable_blocks"] += 1
            status = str(block.get("status") or "").strip()
            if (
                status == "translated"
                and translated_text
                and (contains_cjk_text(translated_text) or is_allowed_untranslated_block(block, translated_text))
            ):
                stats["translated_blocks"] += 1
                stats["translated_chars"] += len(translated_text)
            elif status == "failed":
                stats["failed_blocks"] += 1
            else:
                stats["pending_blocks"] += 1
    translatable = stats["translatable_blocks"]
    if translatable > 0:
        stats["coverage_percent"] = round((stats["translated_blocks"] / translatable) * 100)
        stats["failure_percent"] = round((stats["failed_blocks"] / translatable) * 100)
    return stats


def set_full_translation_diagnostics(item: PaperFullTranslation, pages: list[dict]) -> dict:
    diagnostics = build_full_translation_diagnostics(pages)
    parse_summary = dict(getattr(item, "parse_summary", None) or {})
    parse_summary["diagnostics"] = diagnostics
    item.parse_summary = parse_summary
    return diagnostics


def has_valid_translated_text(block: dict) -> bool:
    translated_text = str(block.get("translated_text") or "").strip()
    return bool(
        translated_text
        and (contains_cjk_text(translated_text) or is_allowed_untranslated_block(block, translated_text))
    )


def reset_translation_pages_for_retry(pages: list[dict], *, failed_only: bool = False) -> int:
    total_units = 0
    for page in pages or []:
        for block in page.get("blocks") or []:
            source_text = str(block.get("source_text") or "").strip()
            if should_copy_block(block) or not source_text:
                block["translate_policy"] = "copy"
                block["status"] = "copied"
                block["translated_text"] = block.get("translated_text") or block.get("source_text", "")
                continue

            if failed_only and has_valid_translated_text(block):
                block["translate_policy"] = "translate"
                block["status"] = "translated"
                continue

            status = str(block.get("status") or "").strip()
            if failed_only and status not in {"failed", "pending"} and has_valid_translated_text(block):
                continue

            block["translate_policy"] = "translate"
            block["status"] = "pending"
            block["translated_text"] = ""
            total_units += 1
    return total_units


def get_translation_source_hash(pages: list[dict]) -> str:
    payload = []
    for page in pages or []:
        payload.append({
            "page_number": page.get("page_number"),
            "width": page.get("width"),
            "height": page.get("height"),
            "layout": page.get("layout") or {},
            "blocks": [
                {
                    "id": block.get("id"),
                    "source_text": block.get("source_text"),
                    "bbox": block.get("bbox"),
                    "column": block.get("column"),
                    "skip_translate": block.get("skip_translate"),
                }
                for block in page.get("blocks") or []
            ],
        })
    return sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def build_translation_cache_key(parse_mode: str, parse_engine: str, local_hash: str, translation_engine: str = "ai", termbase_version: str = "") -> str:
    return f"{parse_mode}:{parse_engine}:{translation_engine}:{termbase_version}:{local_hash}"


def is_completed_translation_cache_hit(
    item: PaperFullTranslation | None,
    *,
    parse_mode: str,
    local_hash: str,
    translation_engine: str = "ai",
    termbase_version: str = "",
) -> bool:
    if not item or item.status != "completed":
        return False
    if any(
        block.get("status") in {"pending", "failed"}
        for page in (item.pages_json or [])
        for block in (page.get("blocks") or [])
    ):
        return False
    source_hash = item.source_hash or ""
    if source_hash == local_hash:
        return True
    if source_hash == build_translation_cache_key(parse_mode, "local", local_hash, translation_engine, termbase_version):
        return True
    legacy_prefix = f"{parse_mode}:local:"
    if source_hash == f"{legacy_prefix}{local_hash}":
        return True
    if parse_mode == "auto" and source_hash.startswith(f"auto:aliyun:{translation_engine}:{termbase_version}:"):
        return True
    if parse_mode == "aliyun" and source_hash.startswith(f"aliyun:aliyun:{translation_engine}:{termbase_version}:"):
        return True
    return False


def load_active_provider(db: Session, user_id: int, provider_id: int | None = None) -> AiProvider | None:
    return resolve_user_provider(
        db,
        user_id,
        provider_id,
        require_active=True,
        fallback_to_active=True,
    )


def contains_cjk_text(text: str) -> bool:
    return any("\u3400" <= char <= "\u9fff" for char in text or "")


def is_allowed_untranslated_block(block: dict, text: str) -> bool:
    import re

    value = str(text or "").strip()
    if block.get("skip_translate"):
        return True
    if not value:
        return False
    if re.match(r"^(https?://|doi:|www\.)", value, re.I):
        return True
    if re.match(r"^[\d\s()[\].,;:/\\+\-=<>%°]+$", value):
        return True
    return False


def get_download_translation_text(block: dict) -> str:
    translated = str(block.get("translated_text") or "").strip()
    source = str(block.get("source_text") or "").strip()
    if block.get("status") == "failed":
        return f"[未译] {source}" if source else ""
    if translated and (contains_cjk_text(translated) or is_allowed_untranslated_block(block, translated)):
        return translated
    if source and is_allowed_untranslated_block(block, source):
        return source
    if translated:
        return f"[未译] {translated}"
    if source:
        return f"[未译] {source}"
    return ""


def _json_list(raw: str | None) -> list:
    if not raw:
        return []
    try:
        value = json.loads(raw)
    except json.JSONDecodeError:
        return []
    return value if isinstance(value, list) else []


def _download_base_name(paper: Paper, fallback: str = "paper") -> str:
    value = str(paper.title or paper.file_name or fallback)
    value = re.sub(r"\.pdf$", "", value, flags=re.I)
    value = re.sub(r'[\\/:*?"<>|]+', " ", value)
    value = re.sub(r"\s+", " ", value).strip()
    return value[:90] or fallback


def _attachment_headers(filename: str) -> dict[str, str]:
    ascii_name = re.sub(r"[^A-Za-z0-9._ -]+", "_", filename).strip(" .") or "download"
    return {
        "Content-Disposition": (
            f'attachment; filename="{ascii_name[:120]}"; '
            f"filename*=UTF-8''{quote(filename)}"
        )
    }


def _hex_to_rgb(color: str | None, fallback: str = "#F3B300") -> tuple[float, float, float]:
    value = (color or fallback).strip()
    if not re.fullmatch(r"#?[0-9A-Fa-f]{3}([0-9A-Fa-f]{3})?", value):
        value = fallback
    value = value.lstrip("#")
    if len(value) == 3:
        value = "".join(ch * 2 for ch in value)
    return (
        int(value[0:2], 16) / 255,
        int(value[2:4], 16) / 255,
        int(value[4:6], 16) / 255,
    )


def _normalized_pdf_rect(page, rect_data: dict):
    import fitz

    page_rect = page.rect
    left = max(0.0, min(1.0, float(rect_data.get("left", 0) or 0)))
    top = max(0.0, min(1.0, float(rect_data.get("top", 0) or 0)))
    width = max(0.0, min(1.0 - left, float(rect_data.get("width", 0) or 0)))
    height = max(0.0, min(1.0 - top, float(rect_data.get("height", 0) or 0)))
    if width <= 0 or height <= 0:
        return None
    x0 = page_rect.x0 + left * page_rect.width
    y0 = page_rect.y0 + top * page_rect.height
    x1 = page_rect.x0 + (left + width) * page_rect.width
    y1 = page_rect.y0 + (top + height) * page_rect.height
    return fitz.Rect(x0, y0, x1, y1)


def _draw_wavy_line(page, rect, color: tuple[float, float, float], width: float) -> None:
    if rect.width <= 1:
        return
    amplitude = max(0.8, min(2.4, rect.height * 0.16))
    step = max(3.0, amplitude * 2.2)
    y = rect.y1 - max(1.2, rect.height * 0.16)
    points = []
    x = rect.x0
    index = 0
    while x <= rect.x1:
        points.append((x, y + (amplitude if index % 2 else -amplitude)))
        x += step
        index += 1
    if points and points[-1][0] < rect.x1:
        points.append((rect.x1, y))
    if len(points) >= 2:
        page.draw_polyline(points, color=color, width=width, overlay=True, stroke_opacity=0.95)


def _render_annotated_pdf(
    file_path: Path,
    annotations: list[Annotation],
    ink_annotations: list[InkAnnotation],
) -> bytes:
    import fitz

    with fitz.open(file_path) as document:
        for annotation in annotations:
            if annotation.page_number < 1 or annotation.page_number > document.page_count:
                continue
            page = document[annotation.page_number - 1]
            color = _hex_to_rgb(annotation.color, "#F3B300")
            for rect_data in _json_list(annotation.rects_json):
                if not isinstance(rect_data, dict):
                    continue
                rect = _normalized_pdf_rect(page, rect_data)
                if not rect:
                    continue
                if annotation.type == "highlight":
                    y_pad = rect.height * 0.16
                    marker_rect = fitz.Rect(rect.x0, rect.y0 + y_pad, rect.x1, rect.y1 - rect.height * 0.08)
                    page.draw_rect(
                        marker_rect,
                        color=None,
                        fill=color,
                        overlay=False,
                        fill_opacity=0.78,
                    )
                elif annotation.type == "wavy_underline":
                    _draw_wavy_line(page, rect, color, max(0.8, min(2.0, rect.height * 0.08)))
                else:
                    y = rect.y1 - max(1.0, rect.height * 0.13)
                    page.draw_line(
                        (rect.x0, y),
                        (rect.x1, y),
                        color=color,
                        width=max(0.8, min(2.0, rect.height * 0.08)),
                        overlay=True,
                        stroke_opacity=0.95,
                    )

        for ink in ink_annotations:
            if ink.page_number < 1 or ink.page_number > document.page_count:
                continue
            page = document[ink.page_number - 1]
            page_rect = page.rect
            points = []
            for point in _json_list(ink.points_json):
                if not isinstance(point, dict):
                    continue
                x = max(0.0, min(1.0, float(point.get("x", 0) or 0)))
                y = max(0.0, min(1.0, float(point.get("y", 0) or 0)))
                points.append((page_rect.x0 + x * page_rect.width, page_rect.y0 + y * page_rect.height))
            if len(points) >= 2:
                page.draw_polyline(
                    points,
                    color=_hex_to_rgb(ink.color, "#15803D"),
                    width=max(0.8, min(18.0, float(ink.stroke_width or 6) * 0.75)),
                    overlay=True,
                    stroke_opacity=max(0.1, min(1.0, float(ink.opacity or 0.85))),
                )

        return document.tobytes(deflate=True, garbage=4)


def _extract_pdf_text_pages(file_path: Path) -> list[tuple[int, str]]:
    try:
        import fitz

        pages: list[tuple[int, str]] = []
        with fitz.open(file_path) as document:
            for index, page in enumerate(document, start=1):
                text = re.sub(r"\s+\n", "\n", page.get_text("text") or "")
                text = re.sub(r"\n{3,}", "\n\n", text).strip()
                pages.append((index, text))
        return pages
    except Exception:
        return []


def _build_annotated_docx(
    paper: Paper,
    file_path: Path,
    annotations: list[Annotation],
    ink_annotations: list[InkAnnotation],
) -> bytes:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Pt, RGBColor

    document = Document()
    document.core_properties.title = paper.title or paper.file_name or "Paper"
    document.add_heading(_download_base_name(paper), 0)

    meta = document.add_paragraph()
    meta.add_run("File: ").bold = True
    meta.add_run(paper.file_name or "")
    if paper.author:
        meta.add_run("\nAuthor: ").bold = True
        meta.add_run(paper.author)
    if paper.doi:
        meta.add_run("\nDOI: ").bold = True
        meta.add_run(paper.doi)
    if paper.arxiv_id:
        meta.add_run("\narXiv: ").bold = True
        meta.add_run(paper.arxiv_id)

    document.add_heading("Paper Text", level=1)
    pages = _extract_pdf_text_pages(file_path)
    if pages:
        for page_number, text in pages:
            document.add_heading(f"Page {page_number}", level=2)
            if not text:
                document.add_paragraph("[No selectable text extracted]")
                continue
            for paragraph in re.split(r"\n{2,}", text):
                clean = re.sub(r"[ \t]+", " ", paragraph).strip()
                if clean:
                    document.add_paragraph(clean)
    else:
        document.add_paragraph("No selectable text could be extracted from this PDF.")

    document.add_page_break()
    document.add_heading("Annotations", level=1)
    if not annotations and not ink_annotations:
        document.add_paragraph("No annotations.")
    type_labels = {
        "highlight": "Highlight",
        "underline": "Underline",
        "wavy_underline": "Wavy underline",
    }
    for annotation in annotations:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"Page {annotation.page_number} - {type_labels.get(annotation.type, annotation.type)}: ").bold = True
        quote_run = paragraph.add_run((annotation.quote_text or "").strip() or "[No quote text]")
        if annotation.type == "highlight":
            quote_run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if annotation.type in {"underline", "wavy_underline"}:
            quote_run.underline = True
        if annotation.color:
            rgb = _hex_to_rgb(annotation.color, "#F3B300")
            quote_run.font.color.rgb = RGBColor(*(round(channel * 255) for channel in rgb))
        quote_run.font.size = Pt(10.5)
    for ink in ink_annotations:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"Page {ink.page_number} - Ink annotation: ").bold = True
        paragraph.add_run(f"{len(_json_list(ink.points_json))} points, color {ink.color or '#15803D'}")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def detect_local_parse_quality(pages: list[dict]) -> dict:
    total_pages = len(pages or [])
    blocks = [block for page in pages or [] for block in page.get("blocks") or []]
    source_texts = [str(block.get("source_text") or "").strip() for block in blocks if str(block.get("source_text") or "").strip()]
    all_text = "\n".join(source_texts)
    text_chars = len(re.sub(r"\s+", "", all_text))
    mojibake_hits = len(re.findall(r"[锟�]|[鐚€畬殏炕璇戞枃鍙傝€冩憳鎽樿]{2,}", all_text))
    tiny_blocks = sum(1 for text in source_texts if len(text) <= 3)
    avg_chars_per_page = text_chars / max(1, total_pages)
    tiny_ratio = tiny_blocks / max(1, len(source_texts))
    reasons = []
    if text_chars < max(80, total_pages * 120):
        reasons.append("文本过少，可能是扫描件或图片型 PDF")
    if avg_chars_per_page < 180:
        reasons.append("平均每页文字过少")
    if mojibake_hits >= 2:
        reasons.append("检测到明显乱码")
    if tiny_ratio > 0.35 and len(source_texts) > 30:
        reasons.append("文本块过碎，版式可能较复杂")
    return {
        "total_pages": total_pages,
        "total_blocks": len(blocks),
        "text_chars": text_chars,
        "avg_chars_per_page": round(avg_chars_per_page, 1),
        "tiny_block_ratio": round(tiny_ratio, 3),
        "mojibake_hits": mojibake_hits,
        "needs_cloud": bool(reasons),
        "reasons": reasons,
    }


def blocks_from_markdown(markdown: str, page_width: float = 595.0, page_height: float = 842.0) -> list[dict]:
    blocks = []
    current: list[str] = []

    def flush() -> None:
        if not current:
            return
        text = " ".join(line.strip() for line in current if line.strip()).strip()
        current.clear()
        if not text:
            return
        index = len(blocks) + 1
        is_heading = text.startswith("#")
        cleaned = re.sub(r"^#+\s*", "", text).strip()
        kind = "heading" if is_heading else "paragraph"
        blocks.append({
            "id": f"p1-a{index}",
            "kind": kind,
            "type": "text",
            "source_text": cleaned,
            "translated_text": "",
            "bbox": [48, 64 + (index - 1) * 28, page_width - 48, 90 + (index - 1) * 28],
            "font_size": 15 if is_heading else 12,
            "font_weight": 700 if is_heading else 400,
            "align": "left",
            "skip_translate": False,
            "translate_policy": "translate",
            "status": "pending",
        })

    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if stripped.startswith(("![](", "![", "<img")):
            flush()
            index = len(blocks) + 1
            blocks.append({
                "id": f"p1-img{index}",
                "kind": "caption",
                "type": "image",
                "source_text": "图片/图表（已保留原文区域）",
                "translated_text": "图片/图表（已保留原文区域）",
                "bbox": [48, 64 + (index - 1) * 28, page_width - 48, 120 + (index - 1) * 28],
                "font_size": 12,
                "font_weight": 400,
                "align": "left",
                "skip_translate": True,
                "translate_policy": "copy",
                "status": "copied",
            })
            continue
        if stripped.startswith("|"):
            flush()
            index = len(blocks) + 1
            blocks.append({
                "id": f"p1-table{index}",
                "kind": "caption",
                "type": "table",
                "source_text": stripped,
                "translated_text": stripped,
                "bbox": [48, 64 + (index - 1) * 28, page_width - 48, 112 + (index - 1) * 28],
                "font_size": 12,
                "font_weight": 400,
                "align": "left",
                "skip_translate": True,
                "translate_policy": "copy",
                "status": "copied",
            })
            continue
        current.append(stripped)
    flush()
    return blocks


def pages_from_aliyun_result(result: dict, fallback_pages: list[dict]) -> list[dict]:
    markdown = str(result.get("markdown") or "").strip()
    if not markdown:
        return fallback_pages
    first_page = fallback_pages[0] if fallback_pages else {}
    width = float(first_page.get("width") or 595)
    height = float(first_page.get("height") or 842)
    blocks = blocks_from_markdown(markdown, width, height)
    return [{
        "page_number": 1,
        "width": width,
        "height": max(height, 120 + len(blocks) * 34),
        "blocks": blocks,
    }]


def maybe_enhance_pages_with_aliyun(paper: Paper, pages: list[dict], parse_mode: str) -> tuple[list[dict], str, dict]:
    quality = detect_local_parse_quality(pages)
    requested = parse_mode == "aliyun"
    forbidden = parse_mode == "local"
    summary = {
        "mode": parse_mode,
        "local_pages": len(pages or []),
        "aliyun_pages": 0,
        "quality": quality,
        "aliyun_available": settings.aliyun_docmind_available,
    }
    if forbidden:
        summary["decision"] = "local_forced"
        return pages, "local", summary
    if not requested and not quality["needs_cloud"]:
        summary["decision"] = "local_quality_ok"
        return pages, "local", summary
    if not settings.aliyun_docmind_available:
        summary["decision"] = "aliyun_unavailable"
        summary["warning"] = "阿里云文档智能未启用，已使用本地解析。"
        return pages, "local", summary

    actual_file = _resolve_paper_file(paper.file_path)
    if not actual_file or not actual_file.exists():
        summary["decision"] = "paper_file_missing"
        summary["warning"] = "论文文件已丢失，无法调用阿里云解析。"
        return pages, "local", summary

    try:
        from app.services.docmind import parse_document_with_aliyun

        result = parse_document_with_aliyun(actual_file, high_precision=requested)
        enhanced_pages = normalize_translation_pages(pages_from_aliyun_result(result, pages))
        summary["decision"] = "aliyun_used"
        summary["aliyun_pages"] = len(enhanced_pages)
        summary["job_id"] = result.get("job_id")
        return enhanced_pages, "aliyun", summary
    except Exception as exc:
        summary["decision"] = "aliyun_failed_fallback_local"
        summary["warning"] = str(exc)[:300]
        return pages, "local", summary


def run_full_translation_task(translation_id: int, provider_id: int | None) -> None:
    from app.db.session import SessionLocal
    from app.services.llm import translate_full_text_blocks
    from app.services.machine_translation import get_translation_engine, translate_with_tencent_mt
    from app.services.termbase import load_termbase

    db = SessionLocal()
    try:
        item = db.get(PaperFullTranslation, translation_id)
        if not item:
            return
        if not settings.full_translation_enabled:
            item.status = "cancelled"
            item.error_message = FULL_TRANSLATION_DISABLED_MESSAGE
            db.add(item)
            db.commit()
            return
        if item.status == "cancelled":
            return

        translation_engine = get_translation_engine()
        paper = db.get(Paper, item.paper_id)
        if not paper:
            item.status = "error"
            item.error_message = "关联论文不存在或已被删除。"
            db.add(item)
            db.commit()
            return

        provider = load_active_provider(db, paper.user_id, provider_id)
        if translation_engine == "ai" and not provider:
            item.status = "error"
            item.error_message = "没有可用的 AI 厂商，请先在 AI 配置中启用一个。"
            db.add(item)
            db.commit()
            if paper:
                create_notification(
                    db,
                    user_id=paper.user_id,
                    source_kind="full_translation",
                    source_id=paper.id,
                    event_kind="failed",
                    title="全文翻译 失败",
                    message=f"{compact_notification_text(paper.title or paper.file_name or '当前论文', 80)} · {compact_notification_text(item.error_message, 120)}",
                    action_kind="open-full-translation",
                    action_payload={"paper_id": paper.id},
                )
            return

        terms, termbase_version = load_termbase()
        item.provider_id = provider.id if provider else None
        item.translation_engine = translation_engine
        item.termbase_version = termbase_version
        item.status = "running"
        item.error_message = None
        db.add(item)
        db.commit()

        api_key = decrypt_api_key(provider.encrypted_api_key) if provider else ""
        pages = deepcopy(item.pages_json or [])
        completed = 0
        cancelled = False
        batch: list[dict[str, str]] = []
        block_refs: list[dict] = []

        def refresh_cancel_state() -> bool:
            nonlocal item
            item = db.get(PaperFullTranslation, translation_id)
            return not item or item.status == "cancelled"

        def flush_batch() -> None:
            nonlocal completed, batch, block_refs, item, pages, cancelled
            if not batch:
                return
            if refresh_cancel_state():
                cancelled = True
                batch = []
                block_refs = []
                return
            translated: dict[str, str] = {}
            if translation_engine == "tencent_mt":
                try:
                    translated = translate_with_tencent_mt(items=batch, terms=terms)
                except Exception:
                    translated = {}
            missing_batch = [entry for entry in batch if entry.get("id") not in translated]
            if missing_batch and provider:
                translated.update(
                    translate_full_text_blocks(
                        base_url=provider.base_url,
                        api_key=api_key,
                        model=provider.model,
                        items=missing_batch,
                    )
                )
            if refresh_cancel_state():
                cancelled = True
                batch = []
                block_refs = []
                return
            for block in block_refs:
                translated_text = str(translated.get(block.get("id")) or "").strip()
                if translated_text and (contains_cjk_text(translated_text) or is_allowed_untranslated_block(block, translated_text)):
                    block["translated_text"] = translated_text
                    block["status"] = "translated"
                    block["translation_engine"] = translation_engine
                else:
                    block["translated_text"] = ""
                    block["status"] = "failed"
                    block["translation_engine"] = translation_engine
            completed += len(batch)
            item = db.get(PaperFullTranslation, translation_id)
            if not item:
                return
            if item.status == "cancelled":
                cancelled = True
                batch = []
                block_refs = []
                return
            item.pages_json = pages
            flag_modified(item, "pages_json")
            set_full_translation_diagnostics(item, pages)
            item.completed_units = completed
            item.status = "running"
            db.add(item)
            db.commit()
            batch = []
            block_refs = []

        for page in pages:
            for block in page.get("blocks") or []:
                if cancelled:
                    break
                source_text = str(block.get("source_text") or "").strip()
                if should_copy_block(block) or not source_text:
                    block["translated_text"] = source_text
                    block["translate_policy"] = "copy"
                    block["status"] = "copied"
                    continue
                if str(block.get("status") or "").strip() == "translated" and has_valid_translated_text(block):
                    continue
                batch.append({"id": block.get("id", ""), "text": source_text})
                block_refs.append(block)
                if len(batch) >= 5:
                    flush_batch()
            if cancelled:
                break
        if not cancelled:
            flush_batch()

        item = db.get(PaperFullTranslation, translation_id)
        if item:
            diagnostics = set_full_translation_diagnostics(item, pages)
            unresolved_blocks = int(diagnostics.get("failed_blocks") or 0) + int(diagnostics.get("pending_blocks") or 0)
            if item.status == "cancelled" or cancelled:
                item.pages_json = pages
                flag_modified(item, "pages_json")
                item.status = "cancelled"
                item.error_message = item.error_message or "已取消全文翻译。"
            elif unresolved_blocks > 0:
                item.pages_json = pages
                flag_modified(item, "pages_json")
                item.completed_units = item.total_units
                item.status = "partial_failed"
                item.error_message = f"全文翻译已生成，但有 {unresolved_blocks} 段未完成，可重新生成。"
            else:
                item.pages_json = pages
                flag_modified(item, "pages_json")
                item.completed_units = item.total_units
                item.status = "completed"
                item.error_message = None
            db.add(item)
            db.commit()
            if item.status in {"completed", "partial_failed"}:
                paper = db.get(Paper, item.paper_id)
                if paper:
                    is_partial = item.status == "partial_failed"
                    create_notification(
                        db,
                        user_id=paper.user_id,
                        source_kind="full_translation",
                        source_id=paper.id,
                        event_kind="failed" if is_partial else "completed",
                        title="全文翻译 部分完成" if is_partial else "全文翻译 已完成",
                        message=(
                            f"{compact_notification_text(paper.title or paper.file_name or '当前论文', 90)} · {unresolved_blocks} 段待处理"
                            if is_partial
                            else compact_notification_text(paper.title or paper.file_name or "当前论文", 120)
                        ),
                        action_kind="open-full-translation",
                        action_payload={"paper_id": paper.id},
                    )
    except Exception as exc:
        item = db.get(PaperFullTranslation, translation_id)
        if item:
            item.status = "error"
            item.error_message = str(exc)[:500]
            db.add(item)
            db.commit()
            paper = db.get(Paper, item.paper_id)
            if paper:
                create_notification(
                    db,
                    user_id=paper.user_id,
                    source_kind="full_translation",
                    source_id=paper.id,
                    event_kind="failed",
                    title="全文翻译 失败",
                    message=f"{compact_notification_text(paper.title or paper.file_name or '当前论文', 80)} · {compact_notification_text(item.error_message, 120)}",
                    action_kind="open-full-translation",
                    action_payload={"paper_id": paper.id},
                )
    finally:
        db.close()


# ── Folders ──────────────────────────────────────────────


@router.get("/folders", response_model=list[FolderResponse])
def list_folders(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> list[FolderResponse]:
    folders = db.scalars(
        select(Folder).where(Folder.user_id == current_user.id).order_by(Folder.id)
    ).all()

    # 确保老用户也有"未分类"文件夹
    if not any(f.name == "未分类" for f in folders):
        uncategorized = Folder(user_id=current_user.id, name="未分类")
        db.add(uncategorized)
        db.commit()
        db.refresh(uncategorized)
        folders = [uncategorized] + list(folders)

    return [build_folder_response(f) for f in folders]


@router.post("/folders", response_model=FolderResponse)
def create_folder(
    payload: FolderCreate,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> FolderResponse:
    existing = db.scalar(
        select(Folder).where(
            Folder.user_id == current_user.id,
            Folder.name == payload.name,
        )
    )
    if existing:
        raise HTTPException(status_code=409, detail="已有同名文件夹。")

    folder = Folder(user_id=current_user.id, name=payload.name)
    db.add(folder)
    db.commit()
    db.refresh(folder)
    return build_folder_response(folder)


@router.patch("/folders/{folder_id}", response_model=FolderResponse)
def rename_folder(
    folder_id: int,
    payload: FolderUpdate,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> FolderResponse:
    folder = db.scalar(
        select(Folder).where(Folder.id == folder_id, Folder.user_id == current_user.id)
    )
    if not folder:
        raise HTTPException(status_code=404, detail="文件夹不存在。")

    if folder.name == "未分类":
        raise HTTPException(status_code=403, detail="未分类文件夹不可修改。")

    # 检查同名
    if payload.name != folder.name:
        existing = db.scalar(
            select(Folder).where(
                Folder.user_id == current_user.id,
                Folder.name == payload.name,
            )
        )
        if existing:
            raise HTTPException(status_code=409, detail="已有同名文件夹。")

    folder.name = payload.name
    db.add(folder)
    db.commit()
    db.refresh(folder)
    return build_folder_response(folder)


@router.delete("/folders/{folder_id}", status_code=status.HTTP_204_NO_CONTENT, response_model=None)
def delete_folder(
    folder_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> None:
    folder = db.scalar(
        select(Folder).where(Folder.id == folder_id, Folder.user_id == current_user.id)
    )
    if not folder:
        raise HTTPException(status_code=404, detail="文件夹不存在。")

    if folder.name == "未分类":
        raise HTTPException(status_code=403, detail="未分类文件夹不可删除。")

    # 将该文件夹下的论文移到用户的"未分类"文件夹
    uncategorized = db.scalar(
        select(Folder).where(
            Folder.user_id == current_user.id,
            Folder.name == "未分类",
        )
    )
    if uncategorized:
        db.execute(
            sql_update(Paper).where(Paper.folder_id == folder_id).values(folder_id=uncategorized.id)
        )

    db.delete(folder)
    db.commit()


# ── Papers ───────────────────────────────────────────────


def _clean_external_id(value: str | None) -> str:
    return str(value or "").strip()


LITERATURE_CACHE_TTL = timedelta(days=7)


def _strip_arxiv_version(arxiv_id: str) -> str:
    return re.sub(r"v\d+$", "", _clean_external_id(arxiv_id), flags=re.IGNORECASE)


def _clean_title_lookup(value: str | None) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())[:300]


def _literature_cache_key(doi: str = "", arxiv_id: str = "", title: str = "") -> str:
    doi = _clean_external_id(doi).lower()
    arxiv_id = _clean_external_id(arxiv_id).lower()
    title = _clean_title_lookup(title).lower()
    if doi:
        return f"doi:{doi}"
    if arxiv_id:
        return f"arxiv:{_strip_arxiv_version(arxiv_id).lower() or arxiv_id}"
    if title:
        return f"title:{sha256(title.encode('utf-8')).hexdigest()}"
    return ""


def _cache_is_fresh(item: PaperLiteratureCache) -> bool:
    updated_at = item.updated_at or item.created_at
    if not updated_at:
        return False
    if updated_at.tzinfo is not None:
        updated_at = updated_at.astimezone(timezone.utc).replace(tzinfo=None)
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    return now - updated_at <= LITERATURE_CACHE_TTL


def _read_literature_cache(db: Session, result_kind: str, lookup_key: str) -> dict | None:
    if not lookup_key:
        return None
    item = db.scalar(
        select(PaperLiteratureCache).where(
            PaperLiteratureCache.result_kind == result_kind,
            PaperLiteratureCache.lookup_key == lookup_key,
        )
    )
    if not item or not _cache_is_fresh(item):
        return None
    return {
        result_kind: item.payload_json or [],
        "source": f"{item.source or '缓存'} · 缓存",
    }


def _write_literature_cache(db: Session, result_kind: str, lookup_key: str, payload: list[dict], source: str) -> None:
    if not lookup_key or not payload:
        return
    item = db.scalar(
        select(PaperLiteratureCache).where(
            PaperLiteratureCache.result_kind == result_kind,
            PaperLiteratureCache.lookup_key == lookup_key,
        )
    )
    if not item:
        item = PaperLiteratureCache(result_kind=result_kind, lookup_key=lookup_key)
    item.payload_json = payload[:1000]
    item.source = str(source or "")[:160]
    db.add(item)
    db.commit()


def _semantic_scholar_paper_keys(doi: str = "", arxiv_id: str = "") -> list[str]:
    keys: list[str] = []
    doi = _clean_external_id(doi)
    arxiv_id = _clean_external_id(arxiv_id)
    if doi:
        keys.append(f"DOI:{doi}")
    if arxiv_id:
        keys.append(f"ARXIV:{arxiv_id}")
        base_arxiv_id = _strip_arxiv_version(arxiv_id)
        if base_arxiv_id and base_arxiv_id != arxiv_id:
            keys.append(f"ARXIV:{base_arxiv_id}")
    return keys


def _semantic_scholar_lookup(doi: str = "", arxiv_id: str = "", fields: str = "paperId") -> dict | None:
    from urllib.request import Request, urlopen

    for paper_key in _semantic_scholar_paper_keys(doi, arxiv_id):
        url = f"https://api.semanticscholar.org/graph/v1/paper/{quote(paper_key, safe=':')}?fields={quote(fields, safe=',')}"
        req = Request(url, headers={"Accept": "application/json"})
        try:
            resp = urlopen(req, timeout=8)
            data = json.loads(resp.read())
            if data.get("paperId"):
                return data
        except Exception:
            continue
    return None


def _format_semantic_scholar_paper(paper: dict) -> dict:
    authors = paper.get("authors") or []
    author_str = ", ".join(a.get("name", "") for a in authors[:3])
    if len(authors) > 3:
        author_str += " et al."
    external_ids = paper.get("externalIds") or {}
    return {
        "title": paper.get("title", ""),
        "authors": author_str,
        "year": paper.get("year"),
        "journal": (paper.get("journal") or {}).get("name") if paper.get("journal") else "",
        "doi": external_ids.get("DOI", ""),
        "arxiv_id": external_ids.get("ArXiv", "") or external_ids.get("ARXIV", ""),
    }


def _crossref_work_year(work: dict) -> int | None:
    for key in ("published-print", "published-online", "published", "issued"):
        parts = (work.get(key) or {}).get("date-parts") or []
        if parts and parts[0]:
            try:
                return int(parts[0][0])
            except Exception:
                continue
    return None


def _format_crossref_work(work: dict) -> dict:
    authors = work.get("author") or []
    return {
        "title": (work.get("title") or [""])[0],
        "authors": "; ".join(
            [item for item in (
                " ".join([str(author.get("given") or "").strip(), str(author.get("family") or "").strip()]).strip()
                for author in authors[:3]
            ) if item]
        ),
        "year": _crossref_work_year(work),
        "journal": (work.get("container-title") or [""])[0],
        "doi": work.get("DOI", ""),
        "arxiv_id": "",
    }


def _lookup_literature_by_title(db: Session, title: str) -> dict | None:
    title = _clean_title_lookup(title)
    if not title:
        return None

    cache_key = _literature_cache_key(title=title)
    cached = _read_literature_cache(db, "lookup", cache_key)
    cached_items = cached.get("lookup") if cached else None
    if cached_items:
        return cached_items[0]

    from urllib.request import Request, urlopen

    def try_crossref_title() -> dict | None:
        url = f"https://api.crossref.org/works?query.title={quote(title)}&rows=1"
        req = Request(url, headers={"Accept": "application/json"})
        resp = urlopen(req, timeout=8)
        data = json.loads(resp.read())
        items = (data.get("message") or {}).get("items") or []
        if not items:
            return None
        work = items[0]
        doi = _clean_external_id(work.get("DOI"))
        if not doi:
            return None
        return {
            **_format_crossref_work(work),
            "doi": doi,
            "source_note": "Crossref 标题匹配",
        }

    def try_openalex_title() -> dict | None:
        url = f"https://api.openalex.org/works?search={quote(title)}&per-page=1"
        req = Request(url, headers={"Accept": "application/json"})
        resp = urlopen(req, timeout=8)
        data = json.loads(resp.read())
        results = data.get("results") or []
        if not results:
            return None
        work = results[0]
        doi = _clean_external_id((work.get("doi") or "").replace("https://doi.org/", ""))
        if not doi:
            return None
        return {
            "title": work.get("title", ""),
            "authors": ", ".join(
                (a.get("author") or {}).get("display_name", "")
                for a in (work.get("authorships") or [])[:3]
            ),
            "year": work.get("publication_year"),
            "journal": ((work.get("primary_location") or {}).get("source") or {}).get("display_name", ""),
            "doi": doi,
            "arxiv_id": "",
            "source_note": "OpenAlex 标题匹配",
        }

    for fn in (try_crossref_title, try_openalex_title):
        try:
            result = fn()
            if result:
                _write_literature_cache(db, "lookup", cache_key, [result], result.get("source_note", "标题匹配"))
                return result
        except Exception:
            continue
    return None


def _resolve_literature_lookup(db: Session, doi: str = "", arxiv_id: str = "", title: str = "") -> dict | None:
    doi = _clean_external_id(doi)
    arxiv_id = _clean_external_id(arxiv_id)
    if doi or arxiv_id:
        return {"doi": doi, "arxiv_id": arxiv_id, "source_note": ""}
    return _lookup_literature_by_title(db, title)


def _is_weak_paper_title(title: str | None, file_name: str | None) -> bool:
    cleaned = _clean_title_lookup(title)
    file_title = _clean_title_lookup(re.sub(r"\.pdf$", "", str(file_name or ""), flags=re.IGNORECASE))
    if not cleaned:
        return True
    if file_title and cleaned.lower() == file_title.lower():
        return True
    if re.fullmatch(r"(arxiv[:\s-]*)?\d{4}\.\d{4,5}(v\d+)?", cleaned, flags=re.IGNORECASE):
        return True
    return False


def _paper_to_metadata(paper: Paper) -> PaperMetadata:
    return PaperMetadata(
        title=paper.title or "",
        author=paper.author,
        subject=paper.subject,
        keywords=paper.keywords,
        creator=paper.creator,
        producer=paper.producer,
        creation_date=paper.creation_date,
        modification_date=paper.modification_date,
        doi=paper.doi,
        arxiv_id=paper.arxiv_id,
        page_count=paper.page_count or 0,
    )


def _apply_local_pdf_metadata(paper: Paper, metadata: PaperMetadata) -> bool:
    changed = False

    def assign(attr: str, value: object, *, force: bool = False) -> None:
        nonlocal changed
        text = compact_paper_text(str(value or ""), PAPER_TEXT_LIMITS.get(attr, 300))
        if is_missing_metadata_value(text):
            return
        current = getattr(paper, attr, None)
        if force or is_missing_metadata_value(current):
            if current != text:
                setattr(paper, attr, text)
                changed = True

    assign(
        "title",
        metadata.title,
        force=is_weak_metadata_title(paper.title, paper.file_name)
        and not is_weak_metadata_title(metadata.title, paper.file_name),
    )
    for attr in (
        "author",
        "subject",
        "keywords",
        "creator",
        "producer",
        "creation_date",
        "modification_date",
        "doi",
        "arxiv_id",
    ):
        assign(attr, getattr(metadata, attr, None))

    if metadata.page_count and not (paper.page_count or 0):
        paper.page_count = metadata.page_count
        changed = True

    return changed


def _apply_external_metadata(paper: Paper, lookup: dict) -> bool:
    changed = False

    def assign(attr: str, value: object, *, force: bool = False) -> None:
        nonlocal changed
        text = compact_paper_text(str(value or ""), PAPER_TEXT_LIMITS.get(attr, 300))
        if not text:
            return
        current = str(getattr(paper, attr) or "").strip()
        if force or not current:
            if current != text:
                setattr(paper, attr, text)
                changed = True

    assign("title", lookup.get("title"), force=_is_weak_paper_title(paper.title, paper.file_name))
    assign("author", lookup.get("authors"))
    assign("subject", lookup.get("journal"))
    assign("doi", lookup.get("doi"))
    assign("arxiv_id", lookup.get("arxiv_id"))
    return changed


@router.post("/{paper_id}/metadata/refresh", response_model=PaperResponse)
def refresh_paper_metadata(
    paper_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> PaperResponse:
    paper = db.scalar(active_paper_query(paper_id, current_user.id))
    if not paper:
        raise HTTPException(status_code=404, detail="论文不存在。")

    changed = False
    actual_file = _resolve_paper_file(paper.file_path)
    if actual_file:
        try:
            local_metadata = extract_pdf_metadata(
                actual_file,
                file_name=paper.file_name or "",
                existing=_paper_to_metadata(paper),
            )
            changed = _apply_local_pdf_metadata(paper, local_metadata) or changed
        except Exception:
            pass

    lookup: dict | None = None
    if paper.doi or paper.arxiv_id:
        external = _semantic_scholar_lookup(
            paper.doi or "",
            paper.arxiv_id or "",
            "title,authors,year,journal,externalIds",
        )
        if external:
            lookup = {
                **_format_semantic_scholar_paper(external),
                "source_note": "Semantic Scholar 标识符匹配",
            }
    if not lookup:
        lookup_title = paper.title or re.sub(r"\.pdf$", "", paper.file_name or "", flags=re.IGNORECASE)
        lookup = _lookup_literature_by_title(db, lookup_title)
    if not lookup:
        if changed:
            db.add(paper)
            db.commit()
            db.refresh(paper)
            return build_paper_response(paper)
        raise HTTPException(status_code=404, detail="未能通过标题、DOI 或 arXiv 匹配到外部文献记录。")

    if _apply_external_metadata(paper, lookup):
        changed = True
    if changed:
        db.add(paper)
        db.commit()
        db.refresh(paper)
    return build_paper_response(paper)


@router.get("/references")
def get_references(
    db: Annotated[Session, Depends(get_db)],
    doi: str = "",
    arxiv_id: str = "",
    title: str = "",
):
    """通过多个 API 获取参考文献"""
    doi = _clean_external_id(doi)
    arxiv_id = _clean_external_id(arxiv_id)
    title = _clean_title_lookup(title)
    lookup = _resolve_literature_lookup(db, doi, arxiv_id, title)
    if lookup:
        doi = _clean_external_id(lookup.get("doi"))
        arxiv_id = _clean_external_id(lookup.get("arxiv_id"))
    if not doi and not arxiv_id:
        return {"references": [], "source": "缺少 DOI/arXiv，标题也未能匹配到外部记录" if title else ""}
    cache_key = _literature_cache_key(doi, arxiv_id, title)
    cached = _read_literature_cache(db, "references", cache_key)
    if cached:
        return cached
    source_note = str((lookup or {}).get("source_note") or "")

    from urllib.request import Request, urlopen

    def try_crossref():
        if not doi:
            return None
        url = f"https://api.crossref.org/works/{doi}"
        req = Request(url, headers={"Accept": "application/json"})
        resp = urlopen(req, timeout=8)
        data = __import__("json").loads(resp.read())
        raw = (data.get("message") or {}).get("reference") or []
        if not raw:
            return None
        refs = []
        for r in raw:
            refs.append({
                "title": r.get("article-title") or r.get("unstructured") or "",
                "authors": r.get("author", ""),
                "year": r.get("year"),
                "journal": r.get("journal-title", ""),
                "doi": r.get("DOI", ""),
            })
        return (refs, f"Crossref ({len(raw)} 条)")

    def try_semantic_scholar():
        paper = _semantic_scholar_lookup(doi, arxiv_id, "paperId")
        paper_id = paper.get("paperId") if paper else None
        if not paper_id:
            return None
        url2 = f"https://api.semanticscholar.org/graph/v1/paper/{paper_id}/references?limit=1000&fields=title,authors,year,journal,externalIds"
        req2 = Request(url2, headers={"Accept": "application/json"})
        resp2 = urlopen(req2, timeout=8)
        data = __import__("json").loads(resp2.read())
        raw = data.get("data") or []
        if not raw:
            return None
        refs = []
        for r in raw:
            p = r.get("citedPaper") or {}
            refs.append(_format_semantic_scholar_paper(p))
        return (refs, f"Semantic Scholar ({len(raw)} 条)")

    def try_openalex():
        import ssl
        ctx = ssl._create_unverified_context()
        if not doi:
            return None
        url = f"https://api.openalex.org/works/doi:{doi}"
        req = Request(url, headers={"Accept": "application/json"})
        resp = urlopen(req, timeout=8)
        data = __import__("json").loads(resp.read())
        ref_ids = data.get("referenced_works") or []
        if not ref_ids:
            return None
        batch = ref_ids[:50]
        batch_str = "|".join(b.rsplit("/", 1)[-1] for b in batch)
        url2 = "https://api.openalex.org/works?filter=" + quote("openalex_id:" + batch_str) + "&per_page=100"
        req2 = Request(url2, headers={"Accept": "application/json"})
        resp2 = urlopen(req2, timeout=10)
        data2 = __import__("json").loads(resp2.read())
        raw = data2.get("results") or []
        refs = []
        for r in raw:
            refs.append({
                "title": r.get("title", ""),
                "authors": ", ".join(
                    (a.get("author") or {}).get("display_name", "")
                    for a in (r.get("authorships") or [])[:3]
                ),
                "year": r.get("publication_year"),
                "journal": ((r.get("primary_location") or {}).get("source") or {}).get("display_name", ""),
                "doi": (r.get("doi") or "").replace("https://doi.org/", ""),
            })
        return (refs, f"OpenAlex ({len(raw)} 条)")

    for fn in (try_crossref, try_semantic_scholar, try_openalex):
        try:
            result = fn()
            if result:
                source = " · ".join([part for part in (source_note, result[1]) if part])
                _write_literature_cache(db, "references", cache_key, result[0], source)
                return {"references": result[0], "source": source}
        except Exception:
            continue

    return {"references": [], "source": "所有来源均无数据"}


@router.get("/citations")
def get_citations(
    db: Annotated[Session, Depends(get_db)],
    doi: str = "",
    arxiv_id: str = "",
    title: str = "",
):
    """获取引用该论文的其他论文"""
    doi = _clean_external_id(doi)
    arxiv_id = _clean_external_id(arxiv_id)
    title = _clean_title_lookup(title)
    lookup = _resolve_literature_lookup(db, doi, arxiv_id, title)
    if lookup:
        doi = _clean_external_id(lookup.get("doi"))
        arxiv_id = _clean_external_id(lookup.get("arxiv_id"))
    if not doi and not arxiv_id:
        return {"citations": [], "source": "缺少 DOI/arXiv，标题也未能匹配到外部记录" if title else ""}
    cache_key = _literature_cache_key(doi, arxiv_id, title)
    cached = _read_literature_cache(db, "citations", cache_key)
    if cached:
        return cached
    source_note = str((lookup or {}).get("source_note") or "")

    from urllib.request import Request, urlopen

    def try_openalex():
        import ssl
        ctx = ssl._create_unverified_context()
        if not doi:
            return None
        url = f"https://api.openalex.org/works/doi:{doi}"
        req = Request(url, headers={"Accept": "application/json"})
        resp = urlopen(req, timeout=8)
        data = __import__("json").loads(resp.read())
        count = data.get("cited_by_count", 0)
        oid = data.get("id", "").rsplit("/", 1)[-1]
        if not oid:
            return None
        req2 = Request("https://api.openalex.org/works?filter=" + quote("cites:" + oid) + "&per_page=100", headers={"Accept": "application/json"})
        resp2 = urlopen(req2, timeout=10)
        data2 = __import__("json").loads(resp2.read())
        refs = []
        for r in data2.get("results") or []:
            refs.append({
                "title": r.get("title", ""),
                "authors": ", ".join(
                    a.get("author", {}).get("display_name", "")
                    for a in (r.get("authorships") or [])[:3]
                ),
                "year": r.get("publication_year"),
                "journal": (r.get("primary_location") or {}).get("source", {}).get("display_name", ""),
                "doi": r.get("doi", "").replace("https://doi.org/", ""),
            })
        return (refs, f"OpenAlex (共被引 {count} 次)")

    def try_semantic_scholar():
        data = _semantic_scholar_lookup(doi, arxiv_id, "paperId,citationCount")
        if not data:
            return None
        paper_id = data.get("paperId")
        count = data.get("citationCount", 0)
        if not paper_id:
            return None
        url2 = f"https://api.semanticscholar.org/graph/v1/paper/{paper_id}/citations?limit=100&fields=title,authors,year,journal,externalIds"
        req2 = Request(url2, headers={"Accept": "application/json"})
        resp2 = urlopen(req2, timeout=8)
        raw = __import__("json").loads(resp2.read()).get("data") or []
        if not raw:
            return None
        refs = []
        for r in raw:
            p = r.get("citingPaper") or {}
            refs.append(_format_semantic_scholar_paper(p))
        return (refs, f"Semantic Scholar (共被引 {count} 次)")

    for name, fn in [("S2", try_semantic_scholar), ("OA", try_openalex)]:
        try:
            result = fn()
            if result:
                source = " · ".join([part for part in (source_note, result[1]) if part])
                _write_literature_cache(db, "citations", cache_key, result[0], source)
                return {"citations": result[0], "source": source}
        except Exception:
            continue

    return {"citations": [], "source": "暂无引用数据"}


@router.get("", response_model=list[PaperResponse])
def list_papers(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
    folder_id: int | None = None,
) -> list[PaperResponse]:
    if _purge_expired_trashed_papers(db, current_user.id):
        db.commit()
    query = select(Paper).where(
        Paper.user_id == current_user.id,
        Paper.deleted_at.is_(None),
    )
    if folder_id is not None:
        query = query.where(Paper.folder_id == folder_id)
    papers = db.scalars(query.order_by(Paper.last_viewed_at.desc(), Paper.created_at.desc())).all()
    return [build_paper_response(p) for p in papers]


@router.get("/trash", response_model=list[PaperTrashResponse])
def list_trash(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> list[PaperTrashResponse]:
    _purge_expired_trashed_papers(db, current_user.id)
    db.commit()
    papers = db.scalars(
        select(Paper)
        .where(
            Paper.user_id == current_user.id,
            Paper.deleted_at.is_not(None),
        )
        .order_by(Paper.deleted_at.desc(), Paper.id.desc())
    ).all()
    folder_ids = {paper.deleted_original_folder_id or paper.folder_id for paper in papers if paper.folder_id}
    folders = db.scalars(
        select(Folder).where(Folder.user_id == current_user.id, Folder.id.in_(folder_ids))
    ).all() if folder_ids else []
    folder_names = {folder.id: folder.name for folder in folders}
    return [
        build_trash_response(
            paper,
            folder_names.get(paper.deleted_original_folder_id or paper.folder_id, "未分类"),
        )
        for paper in papers
    ]


@router.post("/trash/{paper_id}/restore", response_model=PaperResponse)
def restore_paper_from_trash(
    paper_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> PaperResponse:
    _purge_expired_trashed_papers(db, current_user.id)
    paper = db.scalar(
        select(Paper).where(
            Paper.id == paper_id,
            Paper.user_id == current_user.id,
            Paper.deleted_at.is_not(None),
        )
    )
    if not paper:
        db.commit()
        raise HTTPException(status_code=404, detail="回收站中没有这篇文献。")

    target_folder_id = paper.deleted_original_folder_id or paper.folder_id
    target_folder = db.scalar(
        select(Folder).where(Folder.id == target_folder_id, Folder.user_id == current_user.id)
    ) if target_folder_id else None
    if not target_folder:
        target_folder_id = _get_uncategorized_id(db, current_user.id)

    paper.folder_id = target_folder_id
    paper.deleted_at = None
    paper.deleted_original_folder_id = None
    db.add(paper)
    db.commit()
    db.refresh(paper)
    return build_paper_response(paper)


@router.delete("/trash/{paper_id}", status_code=status.HTTP_204_NO_CONTENT, response_model=None)
def permanently_delete_trashed_paper(
    paper_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> None:
    paper = db.scalar(
        select(Paper).where(
            Paper.id == paper_id,
            Paper.user_id == current_user.id,
            Paper.deleted_at.is_not(None),
        )
    )
    if not paper:
        raise HTTPException(status_code=404, detail="回收站中没有这篇文献。")
    _permanently_delete_paper(db, paper)
    db.commit()


@router.delete("/trash")
def empty_trash(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> dict[str, int]:
    papers = db.scalars(
        select(Paper).where(
            Paper.user_id == current_user.id,
            Paper.deleted_at.is_not(None),
        )
    ).all()
    for paper in papers:
        _permanently_delete_paper(db, paper)
    db.commit()
    return {"deleted_count": len(papers)}


@router.post("", response_model=PaperResponse)
async def upload_paper(
    request: Request,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
    file: UploadFile = File(...),
    metadata_json: str = Form(""),
    folder_id: int | None = Form(None),
):
    if file.content_type not in ALLOWED_PDF_TYPES:
        raise HTTPException(status_code=400, detail="仅支持 PDF 格式。")

    content_length = request.headers.get("content-length", "").strip()
    if content_length.isdigit():
        max_request_bytes = settings.papers_max_size_bytes + 512 * 1024
        if int(content_length) > max_request_bytes:
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail=f"PDF file exceeds the {settings.papers_max_size_bytes // (1024 * 1024)} MB upload limit.",
            )

    # 解析 metadata（前端传的 JSON 字符串）
    meta = PaperMetadata()
    if metadata_json.strip():
        try:
            import json
            meta_data = json.loads(metadata_json)
            meta = PaperMetadata(**meta_data)
        except Exception:
            pass

    # 确定目标文件夹
    target_folder_id = folder_id if folder_id is not None else _get_uncategorized_id(db, current_user.id)
    if target_folder_id is not None:
        target = db.scalar(
            select(Folder).where(Folder.id == target_folder_id, Folder.user_id == current_user.id)
        )
        if not target:
            target_folder_id = _get_uncategorized_id(db, current_user.id)

    # 保存文件到磁盘
    papers_dir = Path(settings.papers_upload_dir)
    papers_dir.mkdir(parents=True, exist_ok=True)

    suffix = ".pdf"
    file_name_on_disk = f"{current_user.uid}_{time_ns() // 1_000_000}{suffix}"
    file_path = papers_dir / file_name_on_disk
    file_size_bytes = await _persist_uploaded_pdf(
        file,
        file_path,
        max_size_bytes=settings.papers_max_size_bytes,
    )
    if file_size_bytes <= 0:
        raise HTTPException(status_code=400, detail="上传文件为空。")
    mirror_upload_file(file_path, f"papers/{file_name_on_disk}")
    upload_oss_file(
        file_path,
        paper_object_key(f"/uploads/papers/{file_name_on_disk}"),
        content_type="application/pdf",
    )

    file_url = f"/uploads/papers/{file_name_on_disk}"

    try:
        meta = extract_pdf_metadata(file_path, file_name=file.filename or "", existing=meta)
    except Exception:
        pass

    page_count = meta.page_count
    if not page_count:
        try:
            with fitz.open(file_path) as document:
                page_count = document.page_count
        except Exception:
            page_count = 0

    paper = Paper(
        user_id=current_user.id,
        folder_id=target_folder_id,
        file_name=file.filename or "untitled.pdf",
        file_path=file_url,
        file_size=str(file_size_bytes),
        title=meta.title or (file.filename or "").replace(".pdf", ""),
        author=meta.author,
        subject=meta.subject,
        keywords=meta.keywords,
        creator=meta.creator,
        producer=meta.producer,
        creation_date=meta.creation_date,
        modification_date=meta.modification_date,
        doi=meta.doi,
        arxiv_id=meta.arxiv_id,
        page_count=page_count,
        last_viewed_at=datetime.now(timezone.utc),
    )
    try:
        db.add(paper)
        db.commit()
        db.refresh(paper)
    except Exception:
        try:
            file_path.unlink()
            remove_mirrored_upload(f"papers/{file_name_on_disk}")
        except OSError:
            pass
        raise

    # 自动翻译标题为中文
    discipline = current_user.profile.discipline if current_user.profile else ""
    try:
        translated = translate_title(paper.title, discipline)
        _append_translation_debug_log(
            f"title={paper.title}\ndiscipline={discipline}\ntranslated={translated}\n---\n"
        )
    except Exception as e:
        _append_translation_debug_log(f"ERROR: {e}\n---\n")
        translated = None
    if translated:
        paper.translated_title = translated
        db.add(paper)
        db.commit()
        db.refresh(paper)

    return build_paper_response(paper)


@router.get("/{paper_id}/full-translation", response_model=FullTranslationResponse)
def get_full_translation(
    paper_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> FullTranslationResponse:
    ensure_full_translation_enabled()
    paper = db.scalar(active_paper_query(paper_id, current_user.id))
    if not paper:
        raise HTTPException(status_code=404, detail="论文不存在。")
    item = db.scalar(select(PaperFullTranslation).where(PaperFullTranslation.paper_id == paper_id))
    return build_full_translation_response(item)


@router.post("/{paper_id}/full-translation/start", response_model=FullTranslationResponse)
def start_full_translation(
    paper_id: int,
    payload: FullTranslationStartRequest,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> FullTranslationResponse:
    ensure_full_translation_enabled()
    paper = db.scalar(active_paper_query(paper_id, current_user.id))
    if not paper:
        raise HTTPException(status_code=404, detail="论文不存在。")

    from app.services.machine_translation import get_translation_engine
    from app.services.termbase import load_termbase

    translation_engine = get_translation_engine()
    _terms, termbase_version = load_termbase()
    item = db.scalar(select(PaperFullTranslation).where(PaperFullTranslation.paper_id == paper_id))
    local_pages = normalize_translation_pages([page.model_dump() for page in payload.pages])
    local_hash = payload.source_hash or get_translation_source_hash(local_pages)
    if is_completed_translation_cache_hit(
        item,
        parse_mode=payload.parse_mode,
        local_hash=local_hash,
        translation_engine=translation_engine,
        termbase_version=termbase_version,
    ):
        return build_full_translation_response(item)
    if item and item.status == "running" and item.source_hash.endswith(local_hash):
        return build_full_translation_response(item)

    pages, parse_engine, parse_summary = maybe_enhance_pages_with_aliyun(paper, local_pages, payload.parse_mode)
    parsed_hash = get_translation_source_hash(pages)
    source_hash = build_translation_cache_key(
        payload.parse_mode,
        parse_engine,
        local_hash if parse_engine == "local" else parsed_hash,
        translation_engine,
        termbase_version,
    )
    total_units = count_translatable_units(pages)
    if total_units <= 0:
        raise HTTPException(status_code=400, detail="没有可翻译的正文内容。")

    if not item:
        item = PaperFullTranslation(paper_id=paper_id)

    item.provider_id = payload.provider_id
    item.source_hash = source_hash
    item.parse_mode = payload.parse_mode
    item.parse_engine = parse_engine
    item.parse_summary = parse_summary
    item.translation_engine = translation_engine
    item.termbase_version = termbase_version
    item.status = "running"
    item.pages_json = pages
    item.completed_units = 0
    item.total_units = total_units
    item.error_message = None
    set_full_translation_diagnostics(item, pages)
    db.add(item)
    db.commit()
    db.refresh(item)

    try:
        spawn_full_translation_worker_process(item.id, payload.provider_id)
    except Exception as exc:
        item.status = "error"
        item.error_message = f"启动全文翻译 worker 失败：{exc}"
        db.add(item)
        db.commit()
        db.refresh(item)
    return build_full_translation_response(item)


@router.post("/{paper_id}/full-translation/retry", response_model=FullTranslationResponse)
def retry_full_translation(
    paper_id: int,
    payload: FullTranslationStartRequest,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> FullTranslationResponse:
    ensure_full_translation_enabled()
    item = db.scalar(select(PaperFullTranslation).where(PaperFullTranslation.paper_id == paper_id))
    if item and item.status == "partial_failed" and item.pages_json:
        total_units = reset_translation_pages_for_retry(item.pages_json, failed_only=True)
        if total_units <= 0:
            item.status = "completed"
            item.completed_units = item.total_units
            item.error_message = None
            set_full_translation_diagnostics(item, item.pages_json or [])
            flag_modified(item, "pages_json")
            db.add(item)
            db.commit()
            db.refresh(item)
            return build_full_translation_response(item)
        item.provider_id = payload.provider_id
        item.completed_units = 0
        item.total_units = total_units
        item.status = "running"
        item.error_message = None
        set_full_translation_diagnostics(item, item.pages_json or [])
        flag_modified(item, "pages_json")
        db.add(item)
        db.commit()
        db.refresh(item)
        try:
            spawn_full_translation_worker_process(item.id, payload.provider_id)
        except Exception as exc:
            item.status = "error"
            item.error_message = f"启动全文翻译 worker 失败：{exc}"
            db.add(item)
            db.commit()
            db.refresh(item)
        return build_full_translation_response(item)
    if item:
        db.delete(item)
        db.commit()
    return start_full_translation(paper_id, payload, current_user, db)


@router.post("/{paper_id}/full-translation/cancel", response_model=FullTranslationResponse)
def cancel_full_translation(
    paper_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> FullTranslationResponse:
    ensure_full_translation_enabled()
    paper = db.scalar(active_paper_query(paper_id, current_user.id))
    if not paper:
        raise HTTPException(status_code=404, detail="论文不存在。")
    item = db.scalar(select(PaperFullTranslation).where(PaperFullTranslation.paper_id == paper_id))
    if not item:
        return FullTranslationResponse()
    if item.status == "running":
        item.status = "cancelled"
        item.error_message = "已取消全文翻译。"
        db.add(item)
        db.commit()
        db.refresh(item)
    return build_full_translation_response(item)


@router.get("/{paper_id}/full-translation/stream", response_model=FullTranslationResponse)
def stream_full_translation(
    paper_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> FullTranslationResponse:
    ensure_full_translation_enabled()
    # v1 uses lightweight polling with the stream-shaped endpoint name.
    paper = db.scalar(active_paper_query(paper_id, current_user.id))
    if not paper:
        raise HTTPException(status_code=404, detail="论文不存在。")
    item = db.scalar(select(PaperFullTranslation).where(PaperFullTranslation.paper_id == paper_id))
    return build_full_translation_response(item)


@router.get("/{paper_id}/full-translation/download")
def download_full_translation(
    paper_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
):
    ensure_full_translation_enabled()
    paper = db.scalar(active_paper_query(paper_id, current_user.id))
    if not paper:
        raise HTTPException(status_code=404, detail="论文不存在。")
    item = db.scalar(select(PaperFullTranslation).where(PaperFullTranslation.paper_id == paper_id))
    if not item or item.status not in {"completed", "partial_failed"}:
        raise HTTPException(status_code=404, detail="全文翻译尚未完成。")

    chunks: list[str] = []
    for page in item.pages_json or []:
        chunks.append(f"\n\n# 第 {page.get('page_number')} 页\n")
        for block in page.get("blocks") or []:
            text = get_download_translation_text(block)
            if text:
                chunks.append(text)
    content = "\n\n".join(chunks).strip()
    filename = (paper.title or paper.file_name or "translation").replace("/", " ").replace("\\", " ").strip()
    headers = {"Content-Disposition": f'attachment; filename="{filename[:80]}-translation.md"'}
    return PlainTextResponse(content, media_type="text/markdown; charset=utf-8", headers=headers)


@router.get("/{paper_id}", response_model=PaperResponse)
def get_paper(
    paper_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> PaperResponse:
    paper = db.scalar(active_paper_query(paper_id, current_user.id))
    if not paper:
        raise HTTPException(status_code=404, detail="论文不存在。")
    return build_paper_response(paper)


@router.get("/{paper_id}/file")
async def get_paper_file(
    paper_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
):
    paper = db.scalar(active_paper_query(paper_id, current_user.id))
    if not paper:
        raise HTTPException(status_code=404, detail="论文不存在。")

    object_key = paper_object_key(paper.file_path)
    if settings.oss_direct_download_enabled and oss_object_exists(object_key):
        content = download_oss_bytes(object_key)
        if content:
            return Response(
                content=content,
                media_type="application/pdf",
                headers={
                    "Content-Disposition": f"inline; filename*=UTF-8''{quote(paper.file_name)}",
                    "Cache-Control": "private, max-age=300",
                },
            )

    actual_file = _resolve_paper_file(paper.file_path)
    if not actual_file or not actual_file.exists():
        raise HTTPException(status_code=404, detail="论文文件已丢失。")

    from fastapi.responses import FileResponse
    return FileResponse(
        path=str(actual_file),
        filename=paper.file_name,
        media_type="application/pdf",
    )


@router.get("/{paper_id}/pages/{page_number}/image")
def get_paper_page_image(
    paper_id: int,
    page_number: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
    scale: float = 2.0,
):
    paper = db.scalar(active_paper_query(paper_id, current_user.id))
    if not paper:
        raise HTTPException(status_code=404, detail="论文不存在。")

    actual_file = _resolve_paper_file(paper.file_path)
    if not actual_file or not actual_file.exists():
        raise HTTPException(status_code=404, detail="论文文件已丢失。")

    zoom = max(0.5, min(float(scale or 2.0), 3.0))
    object_key = page_image_object_key(paper.file_path, page_number, zoom)
    if settings.oss_page_image_cache_enabled and oss_object_exists(object_key):
        content = download_oss_bytes(object_key)
        if content:
            return Response(
                content=content,
                media_type="image/png",
                headers={
                    "Cache-Control": "private, max-age=86400",
                },
            )

    try:
        with fitz.open(actual_file) as document:
            if page_number < 1 or page_number > document.page_count:
                raise HTTPException(status_code=404, detail="论文页码不存在。")

            page = document[page_number - 1]
            pixmap = page.get_pixmap(
                matrix=fitz.Matrix(zoom, zoom),
                alpha=False,
                annots=False,
            )
            content = pixmap.tobytes("png")
            if settings.oss_page_image_cache_enabled:
                upload_oss_bytes(content, object_key, content_type="image/png")
            return Response(
                content,
                media_type="image/png",
                headers={
                    "Cache-Control": "private, max-age=86400",
                    "X-Paper-Page-Count": str(document.page_count),
                },
            )
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"页面图片渲染失败：{exc}") from exc


@router.get("/{paper_id}/download/pdf")
def download_paper_pdf(
    paper_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
):
    paper = db.scalar(active_paper_query(paper_id, current_user.id))
    if not paper:
        raise HTTPException(status_code=404, detail="Paper not found")

    actual_file = _resolve_paper_file(paper.file_path)
    if not actual_file or not actual_file.exists():
        raise HTTPException(status_code=404, detail="Paper file is missing")

    annotations = db.scalars(
        select(Annotation)
        .where(Annotation.paper_id == paper_id, Annotation.user_id == current_user.id)
        .order_by(Annotation.page_number, Annotation.start_char, Annotation.id)
    ).all()
    ink_annotations = db.scalars(
        select(InkAnnotation)
        .where(InkAnnotation.paper_id == paper_id, InkAnnotation.user_id == current_user.id)
        .order_by(InkAnnotation.page_number, InkAnnotation.id)
    ).all()

    if not annotations and not ink_annotations:
        from fastapi.responses import FileResponse

        return FileResponse(
            path=str(actual_file),
            filename=paper.file_name,
            media_type="application/pdf",
        )

    ensure_notes_export_allowed(db, current_user.id)

    try:
        content = _render_annotated_pdf(actual_file, annotations, ink_annotations)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Failed to export annotated PDF: {exc}") from exc

    filename = f"{_download_base_name(paper)}-annotated.pdf"
    return Response(
        content,
        media_type="application/pdf",
        headers=_attachment_headers(filename),
    )


@router.get("/{paper_id}/download/word")
def download_paper_word(
    paper_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
):
    paper = db.scalar(active_paper_query(paper_id, current_user.id))
    if not paper:
        raise HTTPException(status_code=404, detail="Paper not found")

    actual_file = _resolve_paper_file(paper.file_path)
    if not actual_file or not actual_file.exists():
        raise HTTPException(status_code=404, detail="Paper file is missing")

    annotations = db.scalars(
        select(Annotation)
        .where(Annotation.paper_id == paper_id, Annotation.user_id == current_user.id)
        .order_by(Annotation.page_number, Annotation.start_char, Annotation.id)
    ).all()
    ink_annotations = db.scalars(
        select(InkAnnotation)
        .where(InkAnnotation.paper_id == paper_id, InkAnnotation.user_id == current_user.id)
        .order_by(InkAnnotation.page_number, InkAnnotation.id)
    ).all()

    ensure_notes_export_allowed(db, current_user.id)

    try:
        content = _build_annotated_docx(paper, actual_file, annotations, ink_annotations)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Failed to export Word document: {exc}") from exc

    filename = f"{_download_base_name(paper)}-annotated.docx"
    return Response(
        content,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers=_attachment_headers(filename),
    )


@router.patch("/{paper_id}", response_model=PaperResponse)
def update_paper(
    paper_id: int,
    payload: PaperUpdate,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> PaperResponse:
    paper = db.scalar(active_paper_query(paper_id, current_user.id))
    if not paper:
        raise HTTPException(status_code=404, detail="论文不存在。")

    if payload.folder_id is not None:
        # 验证目标文件夹属于当前用户
        target = db.scalar(
            select(Folder).where(Folder.id == payload.folder_id, Folder.user_id == current_user.id)
        )
        if not target:
            raise HTTPException(status_code=400, detail="目标文件夹不存在。")
        paper.folder_id = payload.folder_id

    if payload.last_viewed_at:
        paper.last_viewed_at = datetime.now(timezone.utc)

    if payload.title is not None:
        paper.title = payload.title
    if payload.translated_title is not None:
        paper.translated_title = payload.translated_title
    if payload.author is not None:
        paper.author = payload.author
    if payload.subject is not None:
        paper.subject = payload.subject
    if payload.keywords is not None:
        paper.keywords = payload.keywords
    if payload.doi is not None:
        paper.doi = payload.doi
    if payload.arxiv_id is not None:
        paper.arxiv_id = payload.arxiv_id
    if payload.page_count is not None:
        paper.page_count = payload.page_count

    db.add(paper)
    db.commit()
    db.refresh(paper)
    return build_paper_response(paper)


@router.delete("/{paper_id}", status_code=status.HTTP_204_NO_CONTENT, response_model=None)
def delete_paper(
    paper_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> None:
    _purge_expired_trashed_papers(db, current_user.id)
    paper = db.scalar(active_paper_query(paper_id, current_user.id))
    if not paper:
        db.commit()
        raise HTTPException(status_code=404, detail="论文不存在。")

    _clear_generated_content(db, paper.id, current_user.id)
    paper.deleted_at = datetime.now(timezone.utc)
    paper.deleted_original_folder_id = paper.folder_id
    db.add(paper)
    db.commit()


# ── References / Citations ────────────────────────────────


# ── Helpers ──────────────────────────────────────────────


def _get_uncategorized_id(db: Session, user_id: int) -> int:
    folder = db.scalar(
        select(Folder).where(Folder.user_id == user_id, Folder.name == "未分类")
    )
    if folder:
        return folder.id
    # 兜底：如果不存在则创建
    new_folder = Folder(user_id=user_id, name="未分类")
    db.add(new_folder)
    db.flush()
    return new_folder.id


def _resolve_paper_file(file_url: str) -> Path | None:
    """从 URL 路径解析出实际的磁盘文件路径"""
    if not file_url:
        return None
    file_name = Path(file_url).name
    if not file_name:
        return None
    candidate = Path(settings.papers_upload_dir) / file_name
    root = Path(settings.papers_upload_dir).resolve()
    try:
        resolved = candidate.resolve()
    except OSError:
        return None
    if root not in resolved.parents:
        return None
    if not resolved.exists():
        download_oss_file(paper_object_key(file_url), resolved)
    return resolved
